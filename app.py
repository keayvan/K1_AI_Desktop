import streamlit as st
from ollama import Client
import os
import subprocess
import re
from datetime import datetime
import json
import psutil
import time
import base64
from pathlib import Path

APP_VERSION = "1.0.0"

st.set_page_config(page_title=f"K1 AI Desktop v{APP_VERSION}", layout="wide")

client = Client(host='localhost:11434')

if 'chat_history' not in st.session_state:
    st.session_state.chat_history = []

if 'current_folder' not in st.session_state:
    st.session_state.current_folder = ""

DATA_DIR = Path(__file__).parent / "data"
CHATS_DIR = DATA_DIR / "chats"
PROJECTS_FILE = DATA_DIR / "projects.json"
CHATS_DIR.mkdir(parents=True, exist_ok=True)

def load_projects():
    try:
        return json.loads(PROJECTS_FILE.read_text())
    except Exception:
        return {}

def save_projects(projects):
    PROJECTS_FILE.write_text(json.dumps(projects, indent=2, ensure_ascii=False))

def list_chats(project_id):
    chats = []
    for p in CHATS_DIR.glob("*.json"):
        try:
            c = json.loads(p.read_text())
        except Exception:
            continue
        if c.get('project_id') == project_id:
            chats.append(c)
    return sorted(chats, key=lambda c: c.get('updated', ''), reverse=True)

def load_chat(chat_id):
    try:
        return json.loads((CHATS_DIR / f"{chat_id}.json").read_text())
    except Exception:
        return None

def save_chat():
    history = st.session_state.chat_history
    if not history:
        return
    if not st.session_state.chat_id:
        st.session_state.chat_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    first = next((m['content'] for m in history if m['role'] == 'user'), 'New chat')
    path = CHATS_DIR / f"{st.session_state.chat_id}.json"
    old = load_chat(st.session_state.chat_id) or {}
    data = {
        'id': st.session_state.chat_id,
        'title': old.get('title') or first.strip().splitlines()[0][:40],
        'project_id': st.session_state.project_id,
        'created': old.get('created') or datetime.now().isoformat(),
        'updated': datetime.now().isoformat(),
        'messages': history,
    }
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False))

def delete_chat(chat_id):
    (CHATS_DIR / f"{chat_id}.json").unlink(missing_ok=True)

def new_chat():
    st.session_state.chat_history = []
    st.session_state.chat_id = None

if 'chat_id' not in st.session_state:
    st.session_state.chat_id = None

if 'project_id' not in st.session_state:
    st.session_state.project_id = None

if 'stop_processing' not in st.session_state:
    st.session_state.stop_processing = False

if 'uploaded_files_content' not in st.session_state:
    st.session_state.uploaded_files_content = {}

SAFE_COMMANDS = ['ls', 'find', 'grep', 'cat', 'head', 'tail', 'wc', 'git', 'python', 'npm', 'make', 'pytest', 'pip', 'tree', 'stat', 'file']

def process_uploaded_file(uploaded_file):
    file_name = uploaded_file.name
    file_ext = Path(file_name).suffix.lower()
    
    try:
        if file_ext in ['.txt', '.md', '.log']:
            content = uploaded_file.read().decode('utf-8', errors='ignore')
            return {'name': file_name, 'type': 'text', 'content': content[:5000]}
        
        elif file_ext == '.json':
            content = uploaded_file.read().decode('utf-8')
            return {'name': file_name, 'type': 'json', 'content': content[:5000]}
        
        elif file_ext == '.csv':
            content = uploaded_file.read().decode('utf-8')
            return {'name': file_name, 'type': 'csv', 'content': content[:5000]}
        
        elif file_ext == '.pdf':
            try:
                import PyPDF2
                pdf_reader = PyPDF2.PdfReader(uploaded_file)
                text = ""
                for page in pdf_reader.pages[:5]:
                    text += page.extract_text()
                return {'name': file_name, 'type': 'pdf', 'content': text[:5000]}
            except:
                return {'name': file_name, 'type': 'pdf', 'content': '[PDF requires: pip install PyPDF2]'}
        
        elif file_ext == '.docx':
            try:
                from docx import Document
                doc = Document(uploaded_file)
                text = "\n".join([para.text for para in doc.paragraphs])
                return {'name': file_name, 'type': 'docx', 'content': text[:5000]}
            except:
                return {'name': file_name, 'type': 'docx', 'content': '[DOCX requires: pip install python-docx]'}
        
        elif file_ext in ['.png', '.jpg', '.jpeg', '.gif', '.webp']:
            image_data = base64.b64encode(uploaded_file.read()).decode()
            return {'name': file_name, 'type': 'image', 'content': f'[Image: {file_name}]', 'base64': image_data}
        
        elif file_ext in ['.py', '.js', '.ts', '.java', '.cpp', '.c', '.go', '.rs', '.rb']:
            content = uploaded_file.read().decode('utf-8', errors='ignore')
            return {'name': file_name, 'type': 'code', 'content': content[:5000]}
        
        else:
            return {'name': file_name, 'type': 'unknown', 'content': f'[Unsupported: {file_ext}]'}
    
    except Exception as e:
        return {'name': file_name, 'type': 'error', 'content': f'[Error: {str(e)}]'}

def get_system_stats():
    try:
        cpu_percent = psutil.cpu_percent(interval=0.01)
        memory = psutil.virtual_memory()
        return {
            'cpu': cpu_percent,
            'memory_percent': memory.percent,
            'memory_used': memory.used / (1024**3),
            'memory_total': memory.total / (1024**3),
        }
    except:
        return {'cpu': 0, 'memory_percent': 0, 'memory_used': 0, 'memory_total': 0}

def display_system_monitor(elapsed):
    stats = get_system_stats()
    col1, col2, col3, col4, col5 = st.columns(5)
    
    with col1:
        st.metric("CPU", f"{stats['cpu']:.1f}%", delta=None, delta_color="off")
    with col2:
        st.metric("Memory", f"{stats['memory_percent']:.1f}%", delta=None, delta_color="off")
    with col3:
        st.metric("Used", f"{stats['memory_used']:.2f}GB")
    with col4:
        st.metric("Total", f"{stats['memory_total']:.2f}GB")
    with col5:
        st.metric("Time", f"{elapsed:.1f}s")

def read_file(filepath):
    if st.session_state.stop_processing:
        return "[Stopped]"
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    except:
        return "[Cannot read]"

def execute_command(cmd, cwd=None):
    if st.session_state.stop_processing:
        return "[Stopped]"
    try:
        cmd_name = cmd.split()[0]
        if cmd.startswith('git'):
            cmd_name = 'git'
        if cmd_name not in SAFE_COMMANDS:
            return f"[Not allowed: {cmd_name}]"
        result = subprocess.run(cmd, shell=True, cwd=cwd or st.session_state.current_folder, capture_output=True, text=True, timeout=15)
        output = (result.stdout + result.stderr)[:3000]
        return output if output else "[Empty]"
    except subprocess.TimeoutExpired:
        return "[Timeout]"
    except Exception as e:
        return f"[Error: {str(e)}]"

def add_markdown_runs(paragraph, text):
    for i, part in enumerate(re.split(r'\*\*', text)):
        if part:
            paragraph.add_run(part).bold = (i % 2 == 1)

def write_docx(target, content):
    from docx import Document
    doc = Document()
    for line in content.splitlines():
        stripped = line.strip()
        heading = re.match(r'^(#{1,6})\s+(.*)', stripped)
        if heading:
            doc.add_heading(heading.group(2).replace('**', ''), level=min(len(heading.group(1)), 4))
        elif re.match(r'^[-*]\s+', stripped):
            add_markdown_runs(doc.add_paragraph(style='List Bullet'), re.sub(r'^[-*]\s+', '', stripped))
        elif re.match(r'^\d+\.\s+', stripped):
            add_markdown_runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\.\s+', '', stripped))
        elif stripped and set(stripped) != {'-'}:
            add_markdown_runs(doc.add_paragraph(), stripped)
    doc.save(target)

def write_file(rel_path, content):
    if st.session_state.stop_processing:
        return "[Stopped]"
    base = os.path.realpath(st.session_state.current_folder)
    target = os.path.realpath(os.path.join(base, rel_path))
    if os.path.commonpath([base, target]) != base:
        return f"[Not allowed: {rel_path} is outside {base}]"
    try:
        os.makedirs(os.path.dirname(target), exist_ok=True)
        if target.lower().endswith('.docx'):
            write_docx(target, content)
        else:
            with open(target, 'w', encoding='utf-8') as f:
                f.write(content)
        return f"[Wrote {len(content)} chars to {target}]"
    except Exception as e:
        return f"[Error: {str(e)}]"

def scan_folder(folder_path, max_files=30):
    if st.session_state.stop_processing:
        return []
    items = []
    try:
        for root, dirs, files in os.walk(folder_path):
            if st.session_state.stop_processing:
                break
            dirs[:] = [d for d in dirs if not d.startswith('.')]
            for file in sorted(files)[:max_files]:
                if not file.startswith('.'):
                    rel_path = os.path.relpath(os.path.join(root, file), folder_path)
                    items.append(rel_path)
    except:
        pass
    return items

def get_folder_context(folder_path):
    if st.session_state.stop_processing:
        return "[Stopped]"
    context = "Project Files:\n"
    files = scan_folder(folder_path)
    for f in files[:20]:
        context += f"  - {f}\n"
    try:
        branch = subprocess.check_output(['git', '-C', folder_path, 'rev-parse', '--abbrev-ref', 'HEAD'], text=True, stderr=subprocess.DEVNULL, timeout=5).strip()
        context += f"\nBranch: {branch}\n"
    except:
        pass
    return context

def extract_json(text):
    try:
        return json.loads(text)
    except:
        pass
    try:
        start = text.find('{')
        end = text.rfind('}') + 1
        if start >= 0 and end > start:
            json_str = text[start:end]
            return json.loads(json_str)
    except:
        pass
    return {"files": [], "commands": [], "plan": "Parse error"}

with st.sidebar:
    st.write("### Project")
    projects = load_projects()
    options = [None] + list(projects.keys())
    if st.session_state.project_id not in options:
        st.warning("The selected project no longer exists. Switched to chat only.")
        st.session_state.project_id = None
    selected = st.selectbox(
        "Project:", options,
        index=options.index(st.session_state.project_id),
        format_func=lambda pid: "No project (chat only)" if pid is None else projects[pid]['name'],
        label_visibility="collapsed",
    )
    if selected != st.session_state.project_id:
        st.session_state.project_id = selected
        new_chat()
        st.rerun()

    if st.session_state.project_id:
        st.session_state.current_folder = projects[st.session_state.project_id]['folder']
        st.caption(f"Folder: {st.session_state.current_folder}")
        if st.button("Delete project", use_container_width=True):
            for c in list_chats(st.session_state.project_id):
                delete_chat(c['id'])
            del projects[st.session_state.project_id]
            save_projects(projects)
            st.session_state.project_id = None
            new_chat()
            st.rerun()
    else:
        st.session_state.current_folder = ""
        st.caption("No folder: chat only")

    with st.expander("New project"):
        p_name = st.text_input("Name:", key="new_project_name")
        p_folder = st.text_input("Folder:", key="new_project_folder", placeholder="empty = ~/Ollama_projects/<name>")
        if st.button("Create", use_container_width=True):
            if not p_name.strip():
                st.error("Name is required")
            else:
                folder = os.path.expanduser(p_folder.strip()) if p_folder.strip() else os.path.expanduser(f"~/Ollama_projects/{p_name.strip()}")
                try:
                    os.makedirs(folder, exist_ok=True)
                    pid = datetime.now().strftime("%Y%m%d_%H%M%S")
                    projects[pid] = {'name': p_name.strip(), 'folder': os.path.realpath(folder), 'created': datetime.now().isoformat()}
                    save_projects(projects)
                    st.session_state.project_id = pid
                    new_chat()
                    st.rerun()
                except Exception as e:
                    st.error(f"Cannot create folder: {e}")

    st.write("---")
    st.write("### Chats")
    if st.button("+ New chat", use_container_width=True):
        new_chat()
        st.rerun()
    for c in list_chats(st.session_state.project_id):
        col1, col2 = st.columns([5, 1])
        with col1:
            label = ("▶ " if c['id'] == st.session_state.chat_id else "") + c['title']
            if st.button(label, key=f"chat_{c['id']}", use_container_width=True):
                st.session_state.chat_id = c['id']
                st.session_state.chat_history = c.get('messages', [])
                st.rerun()
        with col2:
            if st.button("🗑", key=f"delchat_{c['id']}"):
                delete_chat(c['id'])
                if c['id'] == st.session_state.chat_id:
                    new_chat()
                st.rerun()
    
    st.write("---")
    st.write("### Upload File")
    uploaded_file = st.file_uploader("Choose file:", type=["txt", "md", "json", "csv", "pdf", "docx", "py", "js", "java", "cpp", "png", "jpg", "jpeg"], label_visibility="collapsed")
    
    if uploaded_file:
        file_data = process_uploaded_file(uploaded_file)
        st.session_state.uploaded_files_content[file_data['name']] = file_data
        st.success(f"Loaded: {file_data['name']}")
    
    if st.session_state.uploaded_files_content:
        st.write("**Files:**")
        for fname in st.session_state.uploaded_files_content.keys():
            col1, col2 = st.columns([3, 1])
            with col1:
                st.write(f"- {fname}")
            with col2:
                if st.button("X", key=f"del_{fname}"):
                    del st.session_state.uploaded_files_content[fname]
                    st.rerun()
    
    st.write("---")
    st.write("### Model")
    try:
        installed_models = [m.model for m in client.list().models]
    except Exception:
        installed_models = []
    if not installed_models:
        st.error("No models found. Is Ollama running?")
        installed_models = ["qwen3:8b"]
    default_model = "qwen3:8b" if "qwen3:8b" in installed_models else installed_models[0]
    model = st.selectbox("Select:", installed_models, index=installed_models.index(default_model), label_visibility="collapsed")
    think = st.toggle("Thinking (smarter, slower)", value=False)
    
    st.write("---")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("Clear", use_container_width=True):
            if st.session_state.chat_id:
                delete_chat(st.session_state.chat_id)
            new_chat()
            st.rerun()
    with col2:
        if st.button("Stop", use_container_width=True):
            st.session_state.stop_processing = True
            st.rerun()

st.write(f"## K1 AI Desktop `v{APP_VERSION}`")
st.write("### System Status")
display_system_monitor(0)
st.write("---")

def render_steps(steps):
    for s in steps:
        st.markdown(f"**{s['title']}**")
        if s.get('body'):
            st.code(s['body'], language=s.get('lang'))

def log_step(status, steps, title, body=None, lang=None):
    steps.append({'title': title, 'body': body, 'lang': lang})
    status.update(label=title)
    status.markdown(f"**{title}**")
    if body:
        status.code(body, language=lang)

for msg in st.session_state.chat_history:
    if msg['role'] == 'user':
        st.write(f"**You:** {msg['content']}")
    else:
        if msg.get('steps'):
            with st.expander(f"Steps ({len(msg['steps'])})", expanded=False):
                render_steps(msg['steps'])
        st.write(f"**AI:**\n\n{msg['content']}")
    st.write("---")

user_input = st.chat_input("Message: Enter to send | Shift+Enter for newline")

if user_input:
    st.session_state.stop_processing = False
    st.session_state.chat_history.append({'role': 'user', 'content': user_input})
    save_chat()

    st.write(f"**You:** {user_input}")
    monitor_placeholder = st.empty()
    status = st.status("Starting...", expanded=False)
    response_placeholder = st.empty()

    steps = []
    start_time = time.time()
    last_update_time = start_time

    def refresh_monitor(force=False):
        global last_update_time
        now = time.time()
        if force or now - last_update_time >= 2:
            with monitor_placeholder.container():
                display_system_monitor(now - start_time)
            last_update_time = now

    def check_stop():
        if st.session_state.stop_processing:
            raise KeyboardInterrupt("Stopped")

    try:
        refresh_monitor(force=True)
        check_stop()

        uploaded_context = ""
        if st.session_state.uploaded_files_content:
            uploaded_context = "\n\nUploaded Files:\n"
            for fname, fdata in st.session_state.uploaded_files_content.items():
                uploaded_context += f"\nFile: {fname} ({fdata['type']})\n```\n{fdata['content']}\n```\n"
            log_step(status, steps, "Attached uploaded files",
                     "\n".join(f"- {n} ({d['type']})" for n, d in st.session_state.uploaded_files_content.items()), "text")

        if st.session_state.current_folder:
            folder_context = get_folder_context(st.session_state.current_folder)
            log_step(status, steps, f"1. Scanned folder: {st.session_state.current_folder}", folder_context, "text")


            decision_prompt = f"""{folder_context}{uploaded_context}

    User: {user_input}

    RESPOND WITH ONLY JSON (no other text). Keys:
    - "files": files to read
    - "commands": shell commands to run
    - "write": files to create or overwrite, as {{"path": "...", "content": "..."}} (for .docx, write the content as Markdown; it is converted to a real Word file)
    - "plan": what you will do
    Example:
    {{"files": ["file1.py"], "commands": ["git log -5"], "write": [{{"path": "notes.txt", "content": "text"}}], "plan": "what you will do"}}
    Use empty lists for anything not needed."""

            status.update(label=f"2. Planning with {model}...")
            status.markdown(f"**2. Planning with {model}** (raw model output)")
            plan_box = status.empty()
            full_response = ""
            for chunk in client.generate(model=model, prompt=decision_prompt, stream=True, think=False):
                check_stop()
                full_response += chunk['response']
                plan_box.code(full_response, language="json")
                refresh_monitor()
            steps.append({'title': f"2. Planning with {model} (raw model output)", 'body': full_response, 'lang': "json"})

            decision = extract_json(full_response)
            log_step(status, steps, f"Plan: {decision.get('plan', '')}",
                     json.dumps({'files': decision.get('files', []), 'commands': decision.get('commands', []), 'write': decision.get('write', [])}, indent=2, ensure_ascii=False), "json")

            file_contents = {}
            files_list = decision.get("files", [])[:3]
            for idx, file_path in enumerate(files_list):
                check_stop()
                full_path = os.path.join(st.session_state.current_folder, file_path)
                if os.path.isfile(full_path):
                    content = read_file(full_path)[:2000]
                else:
                    content = "[Not found]"
                file_contents[file_path] = content
                log_step(status, steps, f"3.{idx + 1} Read file: {file_path}", content[:800], "text")
                refresh_monitor()

            command_results = {}
            commands_list = decision.get("commands", [])[:2]
            for idx, cmd in enumerate(commands_list):
                check_stop()
                status.update(label=f"Running: {cmd}")
                result = execute_command(cmd)
                command_results[cmd] = result
                log_step(status, steps, f"4.{idx + 1} Ran command: $ {cmd}", result, "bash")
                refresh_monitor()

            write_results = {}
            write_list = [w for w in decision.get("write", []) if isinstance(w, dict) and w.get("path")][:3]
            for idx, w in enumerate(write_list):
                check_stop()
                result = write_file(w["path"], str(w.get("content", "")))
                write_results[w["path"]] = result
                log_step(status, steps, f"4w.{idx + 1} Wrote file: {w['path']}", f"{result}\n\n{str(w.get('content', ''))[:800]}", "text")
                refresh_monitor()

            check_stop()

            analysis_prompt = f"""{folder_context}{uploaded_context}

    User: {user_input}

    Files: {json.dumps(file_contents, ensure_ascii=False)[:1000]}

    Commands: {json.dumps(command_results, ensure_ascii=False)[:1000]}

    Written files: {json.dumps(write_results, ensure_ascii=False)[:1000]}

    Provide answer."""

            log_step(status, steps, "5. Generating answer...")
            answer = ""
            with response_placeholder.container():
                response_box = st.empty()
                for chunk in client.generate(model=model, prompt=analysis_prompt, stream=True, think=think):
                    check_stop()
                    answer += chunk['response']
                    response_box.markdown(answer)
                    refresh_monitor()
        else:
            log_step(status, steps, "Chat mode (no project folder)")
            messages = []
            if uploaded_context:
                messages.append({'role': 'system', 'content': uploaded_context})
            messages += [{'role': m['role'], 'content': m['content']} for m in st.session_state.chat_history]
            status.update(label=f"Chatting with {model}...")
            answer = ""
            with response_placeholder.container():
                response_box = st.empty()
                for chunk in client.chat(model=model, messages=messages, stream=True, think=think):
                    check_stop()
                    answer += chunk['message']['content']
                    response_box.markdown(answer)
                    refresh_monitor()

        elapsed = time.time() - start_time
        refresh_monitor(force=True)
        steps.append({'title': f"Done in {elapsed:.1f}s", 'body': None, 'lang': None})
        status.update(label=f"Done in {elapsed:.1f}s ({len(steps)} steps)", state="complete", expanded=False)

        st.session_state.chat_history.append({'role': 'assistant', 'content': answer, 'steps': steps})
        save_chat()

        time.sleep(0.5)
        st.rerun()

    except KeyboardInterrupt as e:
        status.update(label=f"Stopped: {e}", state="error")
        st.session_state.stop_processing = False
    except Exception as e:
        status.markdown(f"**Error:** {e}")
        status.update(label=f"Error: {e}", state="error", expanded=True)
